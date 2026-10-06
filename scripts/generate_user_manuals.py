"""Generate MMFC Library user manuals as .docx files."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)
    return h


def add_para(doc, text, bold=False, size=11, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    for run in p.runs:
        set_run_font(run)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(text, style="List Number")
    for run in p.runs:
        set_run_font(run)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, bold=True, size=10)
        # light header shading
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "D6E3F0")
        shading.set(qn("w:val"), "clear")
        hdr[i]._tePr = hdr[i]._tc.get_or_add_tcPr()
        hdr[i]._tc.get_or_add_tcPr().append(shading)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ""
            p = cells[c_idx].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=10)
    doc.add_paragraph()
    return table


def add_callout(doc, title, body):
    p = doc.add_paragraph()
    run = p.add_run(f"{title}: ")
    set_run_font(run, bold=True, size=11, color=RGBColor(0x8B, 0x45, 0x00))
    run2 = p.add_run(body)
    set_run_font(run2, size=11)
    p.paragraph_format.space_after = Pt(10)


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    return doc


def add_title_page(doc, title, subtitle, audience):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("MMFC Library Management System")
    set_run_font(run, size=14, bold=True, color=RGBColor(0x1A, 0x3A, 0x5C))

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p2.add_run(title)
    set_run_font(run, size=22, bold=True, color=RGBColor(0x1A, 0x3A, 0x5C))

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p3.add_run(subtitle)
    set_run_font(run, size=12)

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p4.add_run(audience)
    set_run_font(run, size=11, color=RGBColor(0x55, 0x55, 0x55))

    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p5.add_run("Powered by Pantas · Mindanao Medical Foundation College")
    set_run_font(run, size=10, color=RGBColor(0x77, 0x77, 0x77))

    doc.add_paragraph()


def build_general_manual():
    doc = setup_doc()
    add_title_page(
        doc,
        "Simple User Guide",
        "How to use the library system — written in plain language",
        "For students, employees, library staff, and admins",
    )

    add_heading_styled(doc, "What is this system?", 1)
    add_para(
        doc,
        "This is the college library’s online system. It helps people apply for a library ID, "
        "enter and leave the library by scanning a card or code, look up books and online materials, "
        "and lets library staff approve applications, print IDs, and manage books.",
    )
    add_para(
        doc,
        "Think of it as the library’s digital front desk and attendance desk in one place.",
    )

    add_heading_styled(doc, "Who should read which part?", 1)
    add_table(
        doc,
        ["If you are…", "Go to…"],
        [
            ["A student or employee applying for a library ID", "Part 1"],
            ["Someone using the attendance scanner", "Part 1 — Scanning"],
            ["A student or faculty browsing books online", "Part 1 — Browse books"],
            ["Library staff or an admin", "Part 2"],
        ],
    )

    # Part 1
    add_heading_styled(doc, "Part 1: For Students and Employees", 1)

    add_heading_styled(doc, "A. How to apply for a library ID", 2)
    add_para(doc, "Anyone can apply online. You do not need a login. Library staff will review your application before your ID is ready.")

    add_heading_styled(doc, "Step 1 — Open the registration page", 3)
    add_numbered(doc, "Open the library website.")
    add_numbered(doc, "You should see a Login page that says something like “Welcome! Let’s Begin”.")
    add_numbered(doc, "Click Register.")
    add_para(doc, "(Or ask the library for the registration link.)")

    add_heading_styled(doc, "Step 2 — Choose Student or Employee", 3)
    add_para(doc, "At the top you will see two tabs:")
    add_bullet(doc, "Student — for students")
    add_bullet(doc, "Employee — for faculty / staff / employees")
    add_para(doc, "Click the one that matches you.")

    add_heading_styled(doc, "Step 3 — Fill in your details", 3)
    add_para(doc, "If you are a Student, you will usually need:", bold=True)
    add_bullet(doc, "First name and last name")
    add_bullet(doc, "Student ID (if you have one)")
    add_bullet(doc, "Birth date")
    add_bullet(doc, "Blood type")
    add_bullet(doc, "Course and year level (1st–5th year)")
    add_bullet(doc, "Emergency contact (name, relationship, phone, address)")
    add_bullet(doc, "A 1×1 ID photo (plain white background works best)")
    add_bullet(doc, "Your signature (draw it on the screen; you can clear and try again)")
    add_para(doc, "When finished, click Submit Student Registration.")

    add_para(doc, "If you are an Employee, you will usually need:", bold=True)
    add_bullet(doc, "Name, department, position")
    add_bullet(doc, "Full-time or part-time")
    add_bullet(doc, "Employee ID, birth date, sex, blood type, civil status")
    add_bullet(doc, "Government ID numbers (TIN, PhilHealth, SSS, Pag-IBIG) — numbers only")
    add_bullet(doc, "Emergency contact and address")
    add_bullet(doc, "A formal picture and your signature")
    add_para(doc, "When finished, click Submit Employee Registration.")

    add_heading_styled(doc, "Step 4 — Wait for approval", 3)
    add_para(
        doc,
        "Your application does not create an ID right away. Library staff must review it and Approve it "
        "(or Reject it if something is wrong). After approval, the library can print your ID. "
        "Ask the library when you can pick it up.",
    )

    add_heading_styled(doc, "B. Common words you will see", 2)
    add_table(
        doc,
        ["Word you see", "What it means"],
        [
            ["Patron", "A person registered with the library (student or employee)"],
            ["Pending", "Your application is waiting for staff to review it"],
            ["Approve / Reject", "Staff accepts or declines an application"],
            ["QR Code / RFID", "The code or chip on your ID used for scanning"],
            ["IN / OUT", "You entered or left the library"],
            ["Check Out / Check In", "Borrowing a book or returning a book"],
            ["OPAC / Library Landing", "The book-browsing page for students and faculty"],
        ],
    )

    add_heading_styled(doc, "C. Using the Attendance scanner", 2)
    add_para(doc, "The attendance page is like a kiosk. You usually do not need to log in.")
    add_numbered(doc, "Go to the Attendance screen (or use the computer/tablet set up at the library).")
    add_numbered(doc, "Wait until it says “Ready to scan”.")
    add_numbered(doc, "Scan your library ID (QR code or RFID tag).")
    add_numbered(doc, "The screen will show your name and whether you are IN or OUT.")
    add_callout(
        doc,
        "Tip",
        "If you see “RFID not recognized”, your ID may not be registered yet — ask the library staff.",
    )

    add_heading_styled(doc, "D. Browsing books online", 2)
    add_para(doc, "If the library gave you an account:")
    add_numbered(doc, "Go to the Login page.")
    add_numbered(doc, "Enter your Email and Password, then click Login.")
    add_numbered(doc, "You will land on the Library Landing page.")
    add_para(doc, "From there you can look at New Arrival Books, search by title, filter by course/year, click a book for details, open E-Book materials, and Logout when done.")

    # Part 2
    add_heading_styled(doc, "Part 2: For Library Staff and Admins", 1)
    add_para(
        doc,
        "After you log in as staff or admin, you will usually see the Home (Book Kiosk) page. "
        "The main menu often includes: Home, Attendance, Attendance Logs, ID Generation, and Logout.",
    )

    add_heading_styled(doc, "A. Logging in", 2)
    add_numbered(doc, "Open the library website → Login.")
    add_numbered(doc, "Enter your staff/admin Email and Password.")
    add_numbered(doc, "Click Login. You should arrive at Home.")
    add_para(doc, "When finished for the day, click Logout.")

    add_heading_styled(doc, "B. Approving applications", 2)
    add_numbered(doc, "Click ID Generation.")
    add_numbered(doc, "Click Pending Registrations.")
    add_numbered(doc, "Choose View Students or View Employees.")
    add_numbered(doc, "Review the details and photo.")
    add_numbered(doc, "Click Approve if correct, or Reject if not.")
    add_para(doc, "Approve creates the library record and prepares a QR code so the person can get an ID and use attendance.")

    add_heading_styled(doc, "C. Finding people and printing IDs", 2)
    add_numbered(doc, "Go to ID Generation.")
    add_numbered(doc, "Use the Students or Faculty tabs.")
    add_numbered(doc, "Use Search if needed.")
    add_numbered(doc, "On the person’s row, click Generate, then choose Front, Back, or Download ZIP.")
    add_para(doc, "Admins can also click + Register Patron to add someone without using the pending queue.")

    add_heading_styled(doc, "D. Attendance kiosk", 2)
    add_numbered(doc, "From the menu, open Attendance.")
    add_numbered(doc, "Leave this page open on the kiosk computer.")
    add_numbered(doc, "Ask patrons to scan their ID when entering or leaving.")

    add_heading_styled(doc, "E. Attendance Logs (admin)", 2)
    add_numbered(doc, "Open Attendance Logs.")
    add_numbered(doc, "Set filters (From/To dates, student, course, year) and click Search.")
    add_numbered(doc, "Click Export PDF or Export Excel if you need a file.")

    add_heading_styled(doc, "F. Book borrowing and returns (admin)", 2)
    add_numbered(doc, "Open Logs → Book Check-In & Check-Out Kiosk.")
    add_numbered(doc, "Enter RFID Tag, Patron’s Name, and Action (Check Out or Check In).")
    add_numbered(doc, "Click Record Transaction.")

    add_heading_styled(doc, "G. Adding books and e-resources", 2)
    add_para(doc, "Use Cataloging / Add New Book to add physical books. Use E-Resources Collection to add online materials (journals, e-books, etc.). You can View, Edit, or Delete items later.")

    add_heading_styled(doc, "H. Other admin tools", 2)
    add_bullet(doc, "Prospectus Manager — programs, year levels, and subjects")
    add_bullet(doc, "Repository — upload and store library files")
    add_bullet(doc, "Create Account / User Accounts — login accounts for Student, Staff, or Faculty")

    add_heading_styled(doc, "Quick “What do I do?” cheat sheet", 1)
    add_table(
        doc,
        ["I want to…", "Do this"],
        [
            ["Apply for a library ID", "Login page → Register → Student or Employee → Submit"],
            ["Enter/leave the library", "Use Attendance → scan your ID"],
            ["Browse books with my account", "Login → Library Landing → search or click a book"],
            ["Approve a new registrant", "ID Generation → Pending Registrations → Approve"],
            ["Print an ID", "ID Generation → find person → Generate"],
            ["See who visited", "Attendance Logs → Search → Export if needed"],
            ["Borrow/return a book", "Logs → Check-In & Check-Out → Record Transaction"],
            ["Add a new book", "Cataloging → Add New Book → Save"],
        ],
    )

    add_heading_styled(doc, "Troubleshooting", 1)
    add_table(
        doc,
        ["Problem", "What to try"],
        [
            ["Cannot log in", "Check email/password. Ask staff to reset or create your account."],
            ["“RFID not recognized”", "ID may not be approved yet. Ask library staff."],
            ["Submitted but no ID", "Wait for staff approval, then ask when IDs will be printed."],
            ["Photo/signature won’t upload", "Use a clear photo; try a smaller file; redraw signature."],
            ["Page stuck or blank", "Refresh. Tell staff which button you clicked."],
            ["Wrong Approve/Reject", "Tell an admin right away so they can correct the record."],
        ],
    )

    add_heading_styled(doc, "Tips for less tech-savvy users", 1)
    add_numbered(doc, "Read the button labels — they usually say what will happen (Submit, Approve, Generate, Export).")
    add_numbered(doc, "Do one step at a time — finish a form before opening other tabs.")
    add_numbered(doc, "Save or submit before leaving — closing early may lose your form.")
    add_numbered(doc, "Ask the library if you are unsure — especially before deleting anything.")
    add_numbered(doc, "Always log out on a shared computer when you are done.")

    add_heading_styled(doc, "Need help?", 1)
    add_para(
        doc,
        "Contact the MMFC Library staff for registration problems, ID pickup, login accounts, "
        "or attendance scanner issues.",
    )

    path = OUT / "MMFC_Library_User_Guide.docx"
    doc.save(path)
    return path


def build_registration_manual():
    doc = setup_doc()
    add_title_page(
        doc,
        "Student Registration & QR Guide",
        "How to register a student, approve the application, print the ID, and use the QR code",
        "Step-by-step guide for students and library staff",
    )

    add_heading_styled(doc, "The big picture (in one glance)", 1)
    add_para(
        doc,
        "Here is the full journey from applying to using a library ID:",
    )
    add_numbered(doc, "Student fills out Online Registration.")
    add_numbered(doc, "Application sits in Pending Registrations (waiting).")
    add_numbered(doc, "Staff Approves the application.")
    add_numbered(doc, "System creates a permanent QR code for that student.")
    add_numbered(doc, "Staff prints the ID (Front, Back, or ZIP).")
    add_numbered(doc, "Student scans the QR at Attendance when entering or leaving.")

    add_callout(
        doc,
        "Remember",
        "Submitting a form does NOT print an ID yet. Someone in the library must Approve it first.",
    )

    # Section 1 - Student registers
    add_heading_styled(doc, "1. How a student registers", 1)
    add_para(doc, "No login is needed for this part.")

    add_heading_styled(doc, "Open the form", 2)
    add_numbered(doc, "Go to the library website Login page.")
    add_numbered(doc, "Click Register.")
    add_numbered(doc, "Stay on the Student tab (not Employee).")

    add_heading_styled(doc, "Fill in Personal Information", 2)
    add_table(
        doc,
        ["Field", "Required?", "Simple tip"],
        [
            ["First Name", "Yes", "Include middle initial if you have one (example: JUAN D.)"],
            ["Last Name", "Yes", "Family name"],
            ["Student ID", "No", "Leave blank if you don’t know it. Do NOT type N/A."],
            ["Birth Date", "Yes", "Use your real birthday — not today’s date."],
            ["Blood Type", "No", "Example: O+"],
            ["Course", "Yes", "Your program / course name"],
            ["Year Level", "Yes", "Choose 1ST YEAR through 5TH YEAR"],
        ],
    )

    add_heading_styled(doc, "Fill in Emergency Contact", 2)
    add_table(
        doc,
        ["Field", "Required?", "Simple tip"],
        [
            ["Contact Name", "Yes", "Person to call in an emergency"],
            ["Relationship", "Yes", "Example: Mother, Father, Guardian"],
            ["Contact Number", "Yes", "Working mobile number"],
            ["Address", "No", "Home or contact address"],
        ],
    )

    add_heading_styled(doc, "Photo and signature", 2)
    add_bullet(doc, "Photo: upload a 1×1 ID picture with a plain white background (JPG or PNG). Keep the file under about 4 MB.")
    add_bullet(doc, "Signature: draw with your mouse or finger on the signature box. Use Clear if you want to try again.")
    add_callout(
        doc,
        "Why this matters",
        "Your photo goes on the front of the ID. Your signature and emergency details go on the back.",
    )

    add_heading_styled(doc, "Submit", 2)
    add_numbered(doc, "Double-check names and birth date.")
    add_numbered(doc, "Click Submit Student Registration.")
    add_numbered(doc, "You should see a message that your registration was submitted and is waiting for approval.")
    add_para(doc, "What happens next for the student: wait. Come back or ask the library when your ID is ready to pick up.")

    # Section 2 - Approve
    add_heading_styled(doc, "2. How library staff approves a student", 1)
    add_para(doc, "You must be logged in as Staff or Admin.")

    add_heading_styled(doc, "Find pending applications", 2)
    add_numbered(doc, "Log in to the library system.")
    add_numbered(doc, "Click ID Generation (Registered Students).")
    add_numbered(doc, "Click Pending Registrations.")
    add_numbered(doc, "Click View Students if you are not already on the student list.")

    add_heading_styled(doc, "Review carefully", 2)
    add_para(doc, "Before you Approve, check:")
    add_bullet(doc, "Name spelling")
    add_bullet(doc, "Course and year level")
    add_bullet(doc, "Photo quality (clear face, preferably white background)")
    add_bullet(doc, "Birth date looks real")
    add_bullet(doc, "Emergency contact number looks usable")

    add_heading_styled(doc, "Approve or Reject", 2)
    add_bullet(doc, "Approve — accepts the student into the official registry and creates their permanent QR code.")
    add_bullet(doc, "Reject — removes the application. Use this if the data is wrong or the person should not be registered.")
    add_callout(
        doc,
        "After Approve",
        "The person leaves the Pending list and appears under Registered Students. You can now print their ID.",
    )

    add_heading_styled(doc, "Optional: Admin registers a student directly", 2)
    add_para(
        doc,
        "Admins can skip the waiting list. From ID Generation, click + Register Patron, fill the form, "
        "and Register Patron. The student is created right away (still with a QR code), then you can Generate the ID.",
    )

    # Section 3 - QR
    add_heading_styled(doc, "3. How the QR code works (plain English)", 1)
    add_para(
        doc,
        "A QR code is a square barcode the scanner can read. On the MMFC library ID, the QR is printed on the back of the card.",
    )

    add_heading_styled(doc, "When is the QR created?", 2)
    add_bullet(doc, "When a student first submits online, the system may store a temporary code (it is only for the waiting list).")
    add_bullet(doc, "When staff clicks Approve, the system creates the real QR value used forever for that student.")
    add_bullet(doc, "That real value is a simple number with leading zeros, like 00000001, then 00000002, and so on.")

    add_callout(
        doc,
        "Important",
        "The temporary waiting-list code is NOT what gets printed or scanned. Only the code created after Approve is used on the ID and at Attendance.",
    )

    add_heading_styled(doc, "What is inside the QR?", 2)
    add_para(
        doc,
        "Nothing fancy. The QR simply stores that student number (for example: 00000001). "
        "It does not store the student’s name or a website link. When you scan the card, the system looks up that number in the student list.",
    )

    add_heading_styled(doc, "Where can staff see the QR value?", 2)
    add_para(
        doc,
        "On the Registered Students list, each student shows their QR code as text. "
        "That same value is drawn as a scannable square on the back of the printed ID.",
    )

    add_heading_styled(doc, "How Attendance uses the QR", 2)
    add_numbered(doc, "Open the Attendance page on the kiosk (usually no login needed).")
    add_numbered(doc, "Wait for “Ready to scan”.")
    add_numbered(doc, "Scan the back of the student’s ID (the QR).")
    add_numbered(doc, "The system finds the matching student and marks them IN or OUT.")
    add_para(doc, "First scan of a visit is usually IN. Scanning again when leaving is usually OUT.")
    add_para(
        doc,
        "The same Attendance screen can also recognize a book RFID tag and show whether that book is checked out — "
        "but student attendance and book status are different uses of the same scanner box.",
    )

    add_heading_styled(doc, "If the scan fails", 2)
    add_table(
        doc,
        ["What you see", "Likely reason", "What to do"],
        [
            ["RFID not recognized", "Student not approved yet, or wrong/damaged code", "Confirm the student is in Registered Students; reprint ID if needed"],
            ["Wrong person appears", "QR on card does not match this person", "Do not guess — check the Students list QR text vs the card"],
            ["Nothing happens", "Scanner not focused / page not ready", "Click the Attendance page, wait for Ready to scan, try again"],
        ],
    )

    # Section 4 - Print ID
    add_heading_styled(doc, "4. How to print / download the ID", 1)
    add_para(doc, "Only after the student is approved (or registered by an admin).")

    add_numbered(doc, "Go to ID Generation → Students tab.")
    add_numbered(doc, "Find the student (use Search if the list is long).")
    add_numbered(doc, "Click Generate, then choose one of these:")

    add_table(
        doc,
        ["Option", "What you get"],
        [
            ["Front", "Front of the ID (photo, name, student ID number if any, course)"],
            ["Back", "Back of the ID (QR code, birth date, blood type, emergency info, signature)"],
            ["Download ZIP", "Both front and back images in one zip folder — handy for printing"],
        ],
    )

    add_heading_styled(doc, "What appears on each side", 2)
    add_para(doc, "Front:", bold=True)
    add_bullet(doc, "Student photo")
    add_bullet(doc, "Full name")
    add_bullet(doc, "ID NO. (only if Student ID was filled in)")
    add_bullet(doc, "Course")

    add_para(doc, "Back:", bold=True)
    add_bullet(doc, "QR code (the scannable square)")
    add_bullet(doc, "Birth date and blood type")
    add_bullet(doc, "Emergency contact details")
    add_bullet(doc, "Signature")

    # Section 5 - Checklist
    add_heading_styled(doc, "5. Staff checklist (copy this for training)", 1)
    add_para(doc, "For every new student ID:", bold=True)
    add_numbered(doc, "Student submitted Online Registration (or admin registered them).")
    add_numbered(doc, "Pending application reviewed (name, photo, course, emergency contact).")
    add_numbered(doc, "Clicked Approve.")
    add_numbered(doc, "Confirmed student appears under Registered Students with a QR number.")
    add_numbered(doc, "Generated Front and Back (or Download ZIP).")
    add_numbered(doc, "Printed and cut/laminated the ID according to library practice.")
    add_numbered(doc, "Test-scanned the QR once on the Attendance page before giving the card to the student.")

    # Section 6 - Student tips
    add_heading_styled(doc, "6. Tips for students (share this)", 1)
    add_bullet(doc, "Use a clear 1×1 photo with a white background.")
    add_bullet(doc, "Type your real birth date.")
    add_bullet(doc, "Leave Student ID blank if you are not sure — do not write N/A.")
    add_bullet(doc, "Draw your signature carefully; it will appear on the ID.")
    add_bullet(doc, "After submitting, wait for library approval before expecting an ID.")
    add_bullet(doc, "At the library door, scan the QR on the back of your card when entering and leaving.")
    add_bullet(doc, "Keep the QR clean and uncreased so the scanner can read it.")

    # Section 7 - FAQ
    add_heading_styled(doc, "7. Frequently asked questions", 1)

    add_para(doc, "Q: I registered. Why can’t I scan yet?", bold=True)
    add_para(doc, "A: Staff must Approve you first, then print your ID. Scanning only works with the approved QR.")

    add_para(doc, "Q: Can I register myself as staff without waiting?", bold=True)
    add_para(doc, "A: No. Only an Admin using + Register Patron (or similar admin register) can skip the pending queue.")

    add_para(doc, "Q: Does the QR contain my name or private data?", bold=True)
    add_para(doc, "A: No. It only stores your library QR number. The system looks up your name after the scan.")

    add_para(doc, "Q: What if we Approve the wrong person?", bold=True)
    add_para(doc, "A: Tell an admin immediately. They can edit or remove the student record and reprint if needed.")

    add_para(doc, "Q: Employee / faculty registration — is it the same?", bold=True)
    add_para(
        doc,
        "A: The idea is the same (Register → Pending → Approve → Generate ID → Scan), but the form fields are different "
        "(department, position, government IDs, formal picture). Use the Employee tab and View Employees when approving.",
    )

    add_heading_styled(doc, "Need help?", 1)
    add_para(
        doc,
        "For registration problems, approvals, ID printing, or scanner issues, contact the MMFC Library staff.",
    )

    path = OUT / "MMFC_Student_Registration_and_QR_Guide.docx"
    doc.save(path)
    return path


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    general = build_general_manual()
    reg = build_registration_manual()
    print(f"Wrote: {general}")
    print(f"Wrote: {reg}")


if __name__ == "__main__":
    main()
